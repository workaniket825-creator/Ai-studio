"""Document generation — turn agent content into real binary files in storage.

An LLM tool call can only carry text, so a genuine .docx must be BUILT
server-side from the agent's content and written as bytes. This module does that with
python-docx (permissive OSS) and stores the result in the
configured File Storage via write_bytes, so it lands in the agent's workspace folder and
shows up in the Files screen, downloadable.

Scope: .docx only (headings + paragraphs + optional bullet lists + simple tables). Input is a structure the model
can emit as JSON, or plain markdown-ish text that we lay out reasonably.
"""
from __future__ import annotations

import io

from . import filestore


def _docx_bytes(title: str, content, sections) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    # content can be a plain string (split into paragraphs) …
    if isinstance(content, str) and content.strip():
        for para in content.split("\n\n"):
            p = para.strip()
            if not p:
                continue
            if p.startswith("# "):
                doc.add_heading(p[2:].strip(), level=1)
            elif p.startswith("## "):
                doc.add_heading(p[3:].strip(), level=2)
            elif p.startswith(("- ", "* ")):
                for line in p.splitlines():
                    line = line.strip()
                    if line.startswith(("- ", "* ")):
                        doc.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(p)
    # … or structured sections: [{heading, text?, bullets?, table?}]
    for s in sections or []:
        if s.get("heading"):
            doc.add_heading(s["heading"], level=s.get("level", 1))
        if s.get("text"):
            doc.add_paragraph(s["text"])
        for b in s.get("bullets", []) or []:
            doc.add_paragraph(str(b), style="List Bullet")
        table = s.get("table")
        if table and table.get("rows"):
            rows = table["rows"]
            headers = table.get("headers") or list(rows[0].keys()) if isinstance(rows[0], dict) else None
            if headers:
                t = doc.add_table(rows=1, cols=len(headers))
                t.style = "Light Grid Accent 1"
                for i, h in enumerate(headers):
                    t.rows[0].cells[i].text = str(h)
                for r in rows:
                    cells = t.add_row().cells
                    vals = [r.get(h, "") for h in headers] if isinstance(r, dict) else list(r)
                    for i, v in enumerate(vals):
                        cells[i].text = str(v)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()




def create_document(workspace: str, path: str, fmt: str = "docx",
                    title: str = "", content="", sections=None, sheets=None) -> dict:
    """Render a real .docx and store it in `workspace` (File Storage). Only Word (.docx) is
    supported; any other fmt is rejected (kept minimal + permissively licensed)."""
    fmt = (fmt or "docx").lower().lstrip(".")
    if fmt not in ("docx", "doc"):
        raise filestore.StorageError("only .docx documents are supported")
    data = _docx_bytes(title, content, sections)
    rel = path if path.lower().endswith(".docx") else f"{path}.docx"
    ws = (workspace or "shared").strip("/")
    rel = rel.strip().lstrip("/")
    if not (rel.startswith(ws + "/") or rel == ws):
        rel = f"{ws}/{rel}"
    return filestore.write_bytes(rel, data, actor="agent")
